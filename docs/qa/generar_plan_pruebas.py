"""Plan de casos de prueba del portal de farmacia (LOMS 360).

Genera docs/qa/plan-de-pruebas-farmacia.xlsx: una hoja por rol más una
transversal. Cada caso trae pasos, resultado esperado y columnas para marcar
Pasó / Observación / Quién / Fecha. Las cuentas se nombran por rol (QA Cajero,
etc.); las contraseñas están en qa/accesos_qa.xlsx, que no se versiona.

Regenerar:  python docs/qa/generar_plan_pruebas.py
"""
import pathlib
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = pathlib.Path(__file__).parent
NAVY, AZUL, SUAVE, GRIS = "042C53", "185FA5", "EAF2FB", "E4E7EB"

# (id, pantalla, pasos, resultado esperado)
TRANSVERSAL = [
 ("T01", "Login", "Entrar con una cuenta de QA y su contraseña.", "Abre 'Mis portales' con la tarjeta Farmacia."),
 ("T02", "Login", "Entrar con una contraseña incorrecta.", "Mensaje 'Correo o contraseña incorrectos'. No entra."),
 ("T03", "Olvidé mi contraseña", "Pedir el enlace para QA Cajero (su correo llega a la bandeja del superadmin). Abrirlo y definir una contraseña nueva.", "Llega el correo; el enlace abre la pantalla de nueva contraseña; con la nueva se entra."),
 ("T04", "Cambiar contraseña", "Menú de cuenta → Cambiar contraseña. Primero con la actual mal escrita; después bien.", "Con la actual incorrecta: error. Con la correcta: 'Contraseña actualizada' y se entra con la nueva."),
 ("T05", "Barra del portal", "Dentro de Farmacia, usar 'Cerrar sesión' de la barra. Luego botón Atrás del navegador.", "Vuelve a /login. Atrás no muestra pantallas con datos."),
 ("T06", "Barra del portal", "Dentro de Farmacia, usar 'Mis portales'. Como admin, 'Panel admin'.", "Los botones están siempre visibles a la derecha, aunque la barra de secciones tenga scroll."),
 ("T07", "Moneda", "Recorrer POS, Ventas, Inventario y Finanzas en la Farmacia de Prueba (CLP).", "Todos los montos salen como $ 9.530 (punto de miles, sin decimales). Fechas en formato chileno."),
 ("T08", "Registro", "Con QA Nuevo: 'Crear una cuenta' en /signup, confirmar el correo, iniciar sesión.", "Llega el correo de confirmación; al entrar muestra 'Aprobación pendiente'."),
 ("T09", "Alta por el dueño", "Como dueño: Equipo → Agregar persona → 'Ya tiene cuenta' con el correo de QA Nuevo, rol cajero. Luego entrar con QA Nuevo.", "Queda en el equipo al instante; QA Nuevo entra directo a Farmacia sin esperar aprobación."),
 ("T10", "Contraseña inicial", "Entrar con una cuenta creada desde Equipo ('Crear cuenta nueva'). Intentar ir a /portal/farmacia escribiendo la URL.", "Siempre redirige a 'Elige tu contraseña' hasta que la cambie."),
 ("T11", "Contraseña inicial", "En esa pantalla, poner la inicial y elegir una nueva.", "Entra a Mis portales; en los siguientes ingresos ya no la pide."),
 ("T12", "Admin en 'ver como'", "Como superadmin (o QA Admin): Usuarios → ⋯ → 'Ver sus portales' del dueño de prueba → Farmacia → POS.", "Entra en solo lectura: la barra azul avisa; 'Cobrar' deshabilitado; crear producto no aparece. 'Volver a mi cuenta' sale del modo."),
 ("T13", "Cambio de cuenta", "Cerrar sesión estando en 'ver como' y entrar como QA Cajero en el mismo navegador. Hacer una venta.", "La venta se registra (no aparece 'modo lectura')."),
 ("T14", "Tutorial", "Con cualquier cuenta del portal: 'Ver tutorial' en la barra. Recorrerlo con Siguiente/Anterior y probar 'Ir ahora' en un paso.", "Se abre desde el paso 1, muestra 'Paso N de M' y 'Ir ahora' lleva a la pantalla del paso."),
]

ADMIN = [
 ("A01", "Panel admin · inicio", "Entrar como QA Admin.", "Tarjeta 'Mantenimiento diario' en verde con la última ejecución de hoy. Botón 'Ejecutar ahora' funciona y actualiza la hora."),
 ("A02", "Usuarios · Crear usuario", "Crear un usuario con contraseña generada. Repetir como QA Superadmin.", "QA Admin no ve la opción 'Administrador'; QA Superadmin sí. Las credenciales se muestran una sola vez y se copian."),
 ("A03", "Usuarios · Crear usuario", "Intentar crear con un correo ya existente.", "Mensaje claro: ya existe una cuenta con ese correo."),
 ("A04", "Usuarios", "Aprobar a QA Nuevo (si sigue pendiente), suspenderlo, reactivarlo. Entre cada paso, intentar entrar con QA Nuevo.", "Pendiente → 'Aprobación pendiente'; suspendido → 'Cuenta suspendida'; reactivado → entra."),
 ("A05", "Usuarios · Membresía", "Asignar y extender la membresía del dueño de prueba.", "La fecha de vencimiento cambia; el dueño y su equipo entran al portal."),
 ("A06", "Usuarios · Acceso a portales", "Quitar y volver a dar acceso a Farmacia al dueño de prueba.", "Sin acceso: 'Sin acceso a este portal'. Con acceso: entra."),
 ("A07", "Usuarios · Cambiar rol", "Como QA Admin, buscar la opción de cambiar rol.", "No aparece (solo superadmin)."),
 ("A08", "Auditoría", "Después de A02–A06, abrir Admin → Auditoría.", "Aparecen las acciones con quién, cuándo y sobre quién."),
 ("A09", "Alertas", "Admin → Alertas con una membresía que venza en menos de 7 días.", "La membresía figura en la lista de próximas a vencer."),
 ("A10", "Planes · Límites", "Admin → Planes: editar un plan y poner Máx. regentes 1 y Máx. cajeros 2. Dejar otro plan con los campos vacíos.", "La tarjeta del plan muestra '1 regente · 2 cajeros'; el vacío dice 'sin límite'."),
 ("A11", "Negocios · Nuevo negocio", "Admin → Negocios → Nuevo negocio: 'Farmacia QA Alta', Chile, dueño nuevo (+qa-dueno2 del Gmail), plan con límites, 1 cajero (+qa-cajero3). Dar de alta.", "Muestra las dos contraseñas iniciales una sola vez. La lista suma el negocio con su plan, vencimiento y '1 de 2' cajeros."),
 ("A12", "Negocios · Dueño del alta", "Entrar con el dueño recién creado.", "Pide elegir contraseña; el portal sale en CLP; al crear un producto propone IVA 19; Equipo muestra al cajero."),
 ("A13", "Negocios · Todo o nada", "Repetir el alta usando como dueño 'Ya tiene cuenta' con el correo del dueño de la Farmacia de Prueba.", "Error 'Esa cuenta ya pertenece a otro negocio… Se deshizo lo que se había creado'. En Usuarios no quedó ninguna cuenta nueva."),
 ("A14", "Ver como · Privacidad", "En 'ver como' del dueño de prueba (sin permiso de soporte): abrir Ventas, Finanzas, Recetas, Inventario y Equipo.", "Ventas, Finanzas y Recetas muestran 'sin permiso de soporte'. Inventario y Equipo se ven. El inicio avisa que pacientes y cuentas aparecen en cero."),
 ("A15", "Ver como · Con permiso", "Después del D19 (el dueño da permiso), recargar Ventas en 'ver como'. Luego el dueño quita el permiso y recargar otra vez.", "Con permiso se ven las ventas (sin poder cobrar ni anular). Sin permiso vuelve el aviso. En Auditoría figuran 'autorizar_soporte' y 'revocar_soporte'."),
]

DUENO = [
 ("D01", "Inicio · Alertas de hoy", "Preparar en la Farmacia de Prueba: un lote vencido con stock, uno que vence en 30 días, uno sin fecha, un producto con mínimo 10 y 4 unidades, un tratamiento crónico a 1 día, una cuenta por pagar vencida. Abrir el inicio.", "6 tarjetas con los conteos correctos. Cada una abre la pantalla con el filtro aplicado y el mismo número de filas."),
 ("D02", "Inicio · Alertas de hoy", "Corregir la fecha del lote vencido y volver al inicio.", "La tarjeta de vencidos desaparece al instante. Sin datos de prueba: 'Todo en orden'."),
 ("D03", "Inicio · Datos del negocio", "Botón 'Datos del negocio': cambiar nombre, RUT, dirección y teléfono.", "La etiqueta dice RUT (farmacia chilena). Al guardar, cambia el nombre en la cabecera y la línea gris bajo el nombre."),
 ("D04", "Equipo · Crear cuenta nueva", "Agregar persona → 'Crear cuenta nueva' con nombre, correo (+qa-cajero2 del Gmail), rol cajero, contraseña generada.", "Queda en el equipo con el rol elegido. Se muestran usuario y contraseña una sola vez con botón Copiar."),
 ("D05", "Equipo", "Cambiar el rol de QA Cajero a regente y volver a cajero. Quitar a QA Nuevo del equipo.", "Los cambios se ven al instante; QA Nuevo pierde el acceso al portal."),
 ("D06", "Equipo · Auditoría de accesos", "Mirar la sección de accesos.", "Aparecen los inicios y cierres de sesión recientes de cada miembro."),
 ("D07", "Inventario · Nuevo producto", "Crear un producto con precio 10.000, costo neto 5.000, IVA 19, stock mínimo 5.", "El formulario muestra costo con IVA 5.950, utilidad 4.050 y margen 41 % antes de guardar. En la tabla: Costo c/IVA y Margen."),
 ("D08", "Inventario · Editar", "Editar el precio desde el lápiz de la fila y desde 'Editar producto' en el detalle.", "El cambio se refleja; en el detalle, 'Historial de cambios' muestra precio anterior → nuevo, con quién y cuándo."),
 ("D09", "Inventario · Eliminar", "Eliminar un producto de prueba.", "Desaparece de la tabla; queda en la auditoría."),
 ("D10", "Finanzas", "Abrir Finanzas.", "KPIs del mes, ingresos por método, últimos cierres, 'Inventario valorizado' (a costo, a venta, unidades, por categoría) y 'Día a día'. Los CSV abren en Excel con tildes correctas."),
 ("D11", "Finanzas · Valorización", "Elegir 3 productos con stock, calcular a mano unidades × costo con IVA y comparar con el CSV.", "Coinciden."),
 ("D12", "Compras", "Crear una cuenta por pagar, abonarla parcialmente, anularla.", "Estados pendiente → parcial → anulada. Solo el dueño ve 'Anular'."),
 ("D13", "Cierre", "Abrir Cierre como dueño.", "Ve los totales esperados por método y los movimientos del turno (el cajero no)."),
 ("D14", "Inventario · Mostrar más", "Con los 665 productos cargados, abrir Inventario y escribir rápido en el buscador.", "No se traba. Si hay más de 200 resultados aparece 'Mostrar 200 más'; los contadores de la barra no cambian."),
 ("D15", "Cierre · Retiros y gastos", "Después del C14, abrir Cierre como dueño.", "Los totales del turno muestran 'Retiros: −$ 50.000' y 'Gastos: −$ 5.000'; los movimientos aparecen con signo menos y quién los hizo."),
 ("D16", "Finanzas · Gastos del mes", "Abrir Finanzas y exportar el CSV de 'Gastos y retiros de caja del mes'.", "Gastos por categoría (Insumos $ 5.000) y retiros aparte, con la nota 'no es gasto'. El CSV trae fecha, tipo, categoría, concepto y monto."),
 ("D17", "Datos del negocio · Logo y color", "Subir un logo PNG chico, elegir el color Verde. Después intentar subir un SVG y una imagen de más de 500 KB.", "El logo sale en la barra, el inicio y el comprobante; la pestaña activa queda verde. El SVG y la imagen grande se rechazan con mensaje."),
 ("D18", "Equipo · Suspender", "Con QA Cajero con la sesión abierta en otro navegador: suspenderlo desde ⋯. Que QA Cajero haga clic en cualquier sección. Luego reactivarlo.", "La tarjeta queda 'Suspendido'. QA Cajero sale a Mis portales con 'Acceso suspendido'. Al reactivarlo, vuelve a entrar."),
 ("D19", "Inicio · Soporte de LOMS 360", "En la tarjeta 'Soporte de LOMS 360': 'Permitir por 24 h'. Después 'Quitar permiso'.", "Muestra 'hasta [fecha y hora]' mientras está activo; al quitarlo vuelve a 'tus datos son privados'."),
 ("D20", "Inicio · Primeros pasos", "En un negocio nuevo (el del A11), entrar como dueño y completar cada paso.", "La lista marca cada paso al cumplirse y desaparece cuando están todos."),
 ("D21", "Equipo · Cupo del plan", "En el negocio del A11 (plan con 2 cajeros): agregar un segundo cajero, luego intentar un tercero y cambiar a alguien a cajero.", "Arriba dice '2 de 2 cajeros'; el rol Cajero queda deshabilitado con 'Tu plan no permite más personas con este rol' y en el menú dice '(cupo lleno)'."),
 ("D22", "Equipo · Último acceso", "Mirar las tarjetas del equipo.", "Cada persona muestra 'Último acceso' y, salvo el dueño, 'último cierre' (o 'nunca' / 'ninguno')."),
]

REGENTE = [
 ("R01", "Finanzas", "Entrar como QA Regente a Finanzas.", "Pantalla bloqueada: 'exclusiva del dueño'."),
 ("R02", "Inicio / Equipo", "Buscar el botón 'Datos del negocio' y 'Agregar persona'.", "No aparecen (solo dueño). Sí ve la auditoría de accesos."),
 ("R03", "Inventario · Lotes", "En un producto: ingresar un lote con fecha y otro sin fecha.", "El sin fecha queda gris 'Sin fecha' con botón 'Completar fecha'."),
 ("R04", "Inventario · Editar lote", "Completar la fecha del lote sin fecha. Luego intentar borrarle la fecha a un lote que la tiene.", "La primera se guarda y el semáforo cambia. La segunda se rechaza: 'no puede quedar sin fecha'."),
 ("R05", "Inventario · Movimientos", "Pasar unidades de bodega a venta, registrar una merma con motivo.", "Las cantidades cambian; ambos quedan en 'Movimientos recientes' con el motivo."),
 ("R06", "Inventario · Filtros", "Usar Vencidos, Por vencer, Sin fecha, Bajo mínimo, Sin stock, Inactivos. Probar la URL ?filtro=bajo_minimo y ?filtro=loquesea.", "Cada filtro muestra solo lo que corresponde; la URL abre con el filtro; un valor inválido cae en 'Todos'."),
 ("R07", "Inventario · Stock mínimo", "Poner mínimo 10 a un producto con 4 unidades.", "Chip 'Bajo mínimo' en la fila, píldora en la barra y tarjeta en el inicio. Con mínimo 2 desaparecen."),
 ("R08", "Ventas · Anular", "Anular una venta con motivo.", "Queda marcada ANULADA con el motivo; el stock vuelve al lote del que salió."),
 ("R09", "Pedidos", "Cancelar un pedido con motivo.", "Queda cancelado; si tenía anticipo, aparece el aviso de devolverlo."),
 ("R10", "Compras", "Crear y abonar una cuenta por pagar; buscar 'Anular'.", "Crear y abonar funcionan; 'Anular' no aparece."),
 ("R11", "Inventario · Costos", "Mirar la tabla y el detalle.", "Ve Costo c/IVA, Margen y el historial de cambios."),
 ("R12", "Lotes sin fecha en el POS", "Producto con un lote con fecha y otro sin fecha: venderlo.", "Se descuenta primero el lote con fecha (FEFO); el sin fecha va al final."),
 ("R13", "Cierre · Anular un gasto", "Anular con motivo un gasto registrado por QA Cajero en el turno abierto (la X de la fila). Cerrar caja.", "El gasto queda tachado 'anulado: motivo' y el cierre no lo descuenta."),
 ("R14", "Cierre · Turno cerrado", "Intentar anular un retiro de un turno que ya se cerró (por ejemplo, desde otra pestaña abierta antes del cierre).", "Se rechaza: 'ya entró en un cierre de caja'."),
]

CAJERO = [
 ("C01", "POS · Buscar", "Escribir o escanear un código de barras y Enter. Escribir 'acido' sin tilde.", "El producto entra al ticket al instante. 'acido' encuentra los 'Ácido…'."),
 ("C02", "POS · Sin stock", "Buscar un producto con stock 0 que tenga principio activo.", "Aparece deshabilitado con 'Equivalentes con stock' debajo, si hay otros con el mismo principio activo."),
 ("C03", "POS · Pago mixto", "Vender $ 10.000 pagando $ 6.000 débito + $ 5.000 efectivo.", "Vuelto $ 1.000 (sale del efectivo). Si lo electrónico supera el total, el botón se bloquea con aviso."),
 ("C04", "POS · Receta", "Vender un producto Rx.", "Pide paciente, documento, médico y número de receta; sin eso no deja cobrar. Luego aparece en Recetas."),
 ("C05", "POS · Cliente", "Toca 'Cliente': escribir 1 letra, luego 2; crear un cliente nuevo y usarlo.", "Con 1 letra pide 2; con 2 muestra resultados; el nuevo queda seleccionado sin recargar la página."),
 ("C06", "POS · Después de vender", "Vender 2 unidades de un producto y mirar su stock en la lista. Tocar 'Imprimir' en el aviso verde.", "El stock baja al instante sin recarga. Se abre el comprobante; al imprimir (o 'Guardar como PDF') sale solo el ticket."),
 ("C07", "POS · Dos cajeros", "En dos navegadores, con QA Cajero y el dueño, agregar la última unidad del mismo producto y cobrar en ambos.", "El segundo recibe 'Stock insuficiente' o 'El stock cambió'; su ticket se corrige solo."),
 ("C08", "Ventas", "Abrir Ventas; usar 'Hoy', '7 días', un rango manual; exportar CSV; reimprimir una venta con la impresora.", "La lista y los KPIs cambian con el rango; el CSV abre en Excel; el comprobante se ve igual al del POS. No aparece el botón de anular."),
 ("C09", "Pedidos", "Crear un encargo de $ 10.000 con anticipo de $ 4.000 en efectivo, avanzarlo hasta 'recibido' y usar 'Entregar y cobrar' pagando los $ 6.000 con débito.", "La entrega pasa sin error. En Ventas → Encargos aparecen los dos pagos. En el cierre del turno, lo esperado suma $ 4.000 en efectivo y $ 6.000 en débito."),
 ("C10", "Pacientes", "Crear un tratamiento (30 días), renovarlo, desactivarlo.", "Aparece en la lista con 'Se le acaba'; renovar actualiza la fecha; desactivado queda 'Pausado'."),
 ("C11", "Cierre de caja", "Cerrar caja declarando lo contado por método.", "No se ve el total esperado antes de declarar. Después muestra la diferencia. A los 30 s cierra la sesión (o con 'Salir ahora')."),
 ("C12", "Permisos", "Recorrer Inventario, Finanzas, Compras y Equipo.", "Inventario sin Costo/Margen ni lápiz de editar; Finanzas y Compras bloqueadas; Equipo solo lectura."),
 ("C13", "POS en tablet", "Abrir el POS en una tablet o con el navegador a 800 px de ancho.", "Buscador, resultados y ticket usables sin scroll horizontal."),
 ("C14", "Cierre · Retiro y gasto", "En un turno con una venta de $ 60.000 en efectivo: 'Retiro de caja' de $ 50.000 ('remesa para el dueño') y 'Gasto' de $ 5.000 ('bidón de agua', Insumos). Cerrar declarando $ 5.000 en efectivo.", "Los dos aparecen en Movimientos del turno con signo menos. El cierre da diferencia 0 y dice que ya descuenta los retiros y gastos."),
 ("C15", "Cierre · Anular lo propio", "Registrar un gasto con un monto mal escrito y anularlo con la X, indicando el motivo.", "Queda tachado con 'anulado: motivo'. No aparece la X en movimientos de otras personas."),
 ("C16", "Tutorial · Primer ingreso", "Entrar con una cuenta de cajero nueva (creada en D04), elegir contraseña.", "Se abre solo el tutorial de cajero: vender, cobrar, recetas, encargos, retiros y gastos, cierre. No menciona Finanzas ni Equipo."),
 ("C17", "Tutorial · No se repite", "Tocar 'Saltar tutorial', recargar la página, cerrar sesión y volver a entrar.", "No vuelve a abrirse solo; se puede ver con 'Ver tutorial'."),
]

NUEVO = [
 ("N01", "Registro", "Crear cuenta en /signup con nombre, correo y contraseña de 6 caracteres.", "Pide confirmar el correo; el enlace activa la cuenta."),
 ("N02", "Pendiente", "Entrar antes de que nadie te agregue o apruebe.", "'Aprobación pendiente' con la explicación de los dos caminos."),
 ("N03", "Camino A", "El dueño te agrega desde Equipo con 'Ya tiene cuenta'.", "Entras directo a Farmacia con el rol asignado."),
 ("N04", "Camino B", "Un admin te aprueba, asigna membresía y acceso.", "Entras a Mis portales y a Farmacia como dueño de un negocio nuevo (si un admin lo creó)."),
]

HOJAS = [
 ("Transversal", TRANSVERSAL), ("Admin", ADMIN), ("Dueño", DUENO),
 ("Regente", REGENTE), ("Cajero", CAJERO), ("Usuario nuevo", NUEVO),
]

ENCABEZADOS = ["ID", "Pantalla", "Pasos", "Resultado esperado", "Pasó", "Observación / falla", "Quién probó", "Fecha"]
ANCHOS      = [7, 24, 60, 60, 8, 40, 16, 12]

wb = Workbook()
res = wb.active; res.title = "Cómo usar"
fino = Side(style="thin", color=GRIS)
borde = Border(left=fino, right=fino, top=fino, bottom=fino)

instrucciones = [
 ("Plan de pruebas · Portal de farmacia LOMS 360", "titulo"),
 ("Versión 1.1 · 21 de octubre de 2026", "sub"),
 ("", ""),
 ("Cuentas", "h"),
 ("Usa las cuentas del kit de QA (archivo accesos_qa.xlsx, no versionado): QA Superadmin, QA Admin, Dueño de Prueba (farmacia.prueba@ejemplo.com), QA Regente, QA Cajero y QA Nuevo. Todas trabajan sobre la 'Farmacia de Prueba'. Nunca probar en la Farmacia Luis Gabriel: es el negocio real.", ""),
 ("", ""),
 ("Dónde", "h"),
 ("En la dirección del portal que te pase el equipo (producción o el preview del PR que se esté probando). Los datos de prueba (lotes, productos, tratamientos) se crean en la Farmacia de Prueba y se pueden borrar al terminar.", ""),
 ("", ""),
 ("Cómo marcar", "h"),
 ("Cada hoja es un rol. Por caso: haz los pasos, compara con 'Resultado esperado' y marca Pasó = Sí o No. Si es No, escribe qué pasó (mensaje exacto, pantalla, hora) y, si puedes, agrega una captura con el ID del caso en el nombre del archivo. Quién probó y Fecha ayudan a reproducir.", ""),
 ("", ""),
 ("Novedades de la versión 1.1", "h"),
 ("Lo de la reunión del 21 de octubre: retiros y gastos de caja (C14, C15, R13, R14, D15, D16), tutorial (T14, C16, C17), alta de negocios y límites por plan (A10–A13, D20, D21), logo, color y suspender (D17, D18, D22) y privacidad del soporte (A14, A15, D19). C09 ahora prueba el cobro de encargos con montos concretos. Estos casos requieren que estén desplegados los cambios correspondientes y corridos sus SQL.", ""),
 ("", ""),
 ("Orden sugerido", "h"),
 ("1) Transversal T01–T07 · 2) Admin · 3) Dueño (deja los datos de prueba del D01 creados) · 4) Regente · 5) Cajero · 6) Usuario nuevo · 7) Transversal T08–T13.", ""),
 ("", ""),
 ("Qué NO cubre este plan", "h"),
 ("Boleta electrónica (el POS emite un comprobante interno, no documento tributario), respaldos y cargas masivas: se prueban aparte.", ""),
]
for i, (texto, tipo) in enumerate(instrucciones, 1):
    c = res.cell(row=i, column=1, value=texto)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if tipo == "titulo": c.font = Font(bold=True, size=16, color=NAVY)
    elif tipo == "sub":  c.font = Font(color="6E7681")
    elif tipo == "h":    c.font = Font(bold=True, color=AZUL, size=12)
res.column_dimensions["A"].width = 110

total = 0
for nombre, casos in HOJAS:
    ws = wb.create_sheet(nombre)
    ws.append(ENCABEZADOS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(vertical="center", wrap_text=True); c.border = borde
    for id_, pantalla, pasos, esperado in casos:
        ws.append([id_, pantalla, pasos, esperado, "", "", "", ""])
        total += 1
    for i, w in enumerate(ANCHOS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            c.alignment = Alignment(vertical="top", wrap_text=True); c.border = borde
        fila[0].font = Font(bold=True, color=AZUL)
        fila[4].alignment = Alignment(horizontal="center", vertical="top")
    dv = DataValidation(type="list", formula1='"Sí,No,Parcial"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"E2:E{len(casos) + 1}")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 28

wb.save(AQUI / "plan-de-pruebas-farmacia.xlsx")
print("ok:", total, "casos en", len(HOJAS), "hojas")
